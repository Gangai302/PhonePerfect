import docx

doc = docx.Document("PhonePerfect_Project_Report.docx")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
print(f"Total Sections: {len(doc.sections)}")

# Check for residual references to DocSphere
docsphere_count = 0
for idx, p in enumerate(doc.paragraphs):
    if "docsphere" in p.text.lower() or "gemma" in p.text.lower() or "mediapipe" in p.text.lower():
        print(f"Residual in P{idx}: {p.text[:60]}")
        docsphere_count += 1

for tidx, t in enumerate(doc.tables):
    for r in t.rows:
        for c in r.cells:
            if "docsphere" in c.text.lower() or "gemma" in c.text.lower() or "mediapipe" in c.text.lower():
                print(f"Residual in Table {tidx}: {c.text[:60]}")
                docsphere_count += 1

print(f"Total DocSphere/residual hits found: {docsphere_count}")
