import docx

doc = docx.Document("PhonePerfect_Project_Report.docx")

with open("docsphere_clean_structure.txt", "w", encoding="utf-8") as out:
    for i, p in enumerate(doc.paragraphs):
        text_preview = p.text.replace("\n", " ").strip()[:80]
        out.write(f"P{i:03d} [{p.style.name}] (align={p.alignment}): {text_preview}\n")

    for ti, t in enumerate(doc.tables):
        header_text = " | ".join([c.text.replace("\n", " ").strip() for c in t.rows[0].cells])
        out.write(f"TBL{ti:02d} ({len(t.rows)}x{len(t.columns)}): {header_text[:100]}\n")

print("Wrote docsphere_clean_structure.txt")
