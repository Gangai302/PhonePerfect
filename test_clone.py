import docx
import shutil
import os

TEMPLATE_PATH = r"C:\Users\nisha\Downloads\docsphere.docx"
OUTPUT_PATH = r"c:\Users\nisha\OneDrive\Desktop\scratch\PhonePerfect_Project_Report.docx"

shutil.copyfile(TEMPLATE_PATH, OUTPUT_PATH)
doc = docx.Document(OUTPUT_PATH)
print(f"Loaded cloned document: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.sections)} sections.")
